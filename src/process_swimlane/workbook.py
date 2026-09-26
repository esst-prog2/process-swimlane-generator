"""Read the fixed Process input area without modifying the workbook."""

from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
import re

from openpyxl import load_workbook

from .model import Diagnostic, Task, ValidationError

HEADERS = ("task_id", "section", "task", "department", "annual_frequency", "next_task_id")
DECIMAL_TEXT = re.compile(r"[+]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)\Z")


def effective_value(cell, saved):
    if cell.data_type == "f":
        if saved.value is None or saved.data_type == "e":
            raise ValidationError(Diagnostic(
                "Formula has no usable saved value. Recalculate and save the workbook in Excel before running the generator.",
                cell.row, cell.column_letter,
            ))
        return saved.value
    if cell.data_type == "e":
        raise ValidationError(Diagnostic("Excel error value is not valid input.", cell.row, cell.column_letter))
    return cell.value


def text_value(value, row, column, name, optional=False):
    if optional and value in (None, ""):
        return None
    if not isinstance(value, str):
        message = "Task IDs must be stored as text (format the cell as Text and re-enter the ID)." if name in ("task_id", "next_task_id") else f"{name} must contain nonblank text."
        raise ValidationError(Diagnostic(message, row, column))
    if not value.strip():
        raise ValidationError(Diagnostic(f"{name} must contain nonblank text.", row, column))
    return value


def frequency_value(value, row):
    if value is None or value == "":
        return None
    try:
        if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
            raise ValueError
        if isinstance(value, str) and not DECIMAL_TEXT.fullmatch(value):
            raise ValueError
        number = Decimal(str(value))
        if not number.is_finite() or number < 0:
            raise ValueError
        return number
    except (ValueError, InvalidOperation):
        raise ValidationError(Diagnostic("annual_frequency must be a finite non-negative number using a dot decimal separator, or blank.", row, "E")) from None


def read_workbook(path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValidationError(Diagnostic("Input must be an .xlsx workbook."))
    snapshot = path.read_bytes()
    formulas = load_workbook(BytesIO(snapshot), data_only=False)
    saved = None
    try:
        saved = load_workbook(BytesIO(snapshot), data_only=True)
        if "Process" not in formulas.sheetnames:
            raise ValidationError(Diagnostic("Required worksheet 'Process' is missing."))
        sheet, values = formulas["Process"], saved["Process"]
        for merged in sheet.merged_cells.ranges:
            if merged.min_col <= 6:
                raise ValidationError(Diagnostic(f"Merged input cells {merged} are unsupported; unmerge the Process A-F input area."))
        errors = []
        for index, header in enumerate(HEADERS, 1):
            cell = sheet.cell(1, index)
            if cell.value != header or cell.data_type == "f":
                errors.append(Diagnostic(f"Expected required header {header!r}.", 1, cell.column_letter))
        if errors:
            raise ValidationError(*errors)
        tasks = []
        for cells in sheet.iter_rows(min_row=2, max_col=6):
            if all(cell.value is None for cell in cells):
                continue
            row = cells[0].row
            try:
                data = [effective_value(cell, values.cell(row, cell.column)) for cell in cells]
                tasks.append(Task(
                    row,
                    text_value(data[0], row, "A", "task_id"),
                    text_value(data[1], row, "B", "section"),
                    text_value(data[2], row, "C", "task"),
                    text_value(data[3], row, "D", "department"),
                    frequency_value(data[4], row),
                    text_value(data[5], row, "F", "next_task_id", optional=True),
                ))
            except ValidationError as error:
                errors.extend(error.diagnostics)
        if errors:
            raise ValidationError(*errors)
        return tasks
    finally:
        formulas.close()
        if saved is not None:
            saved.close()
