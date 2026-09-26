from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from decimal import Decimal

from openpyxl import Workbook
from process_swimlane.model import ValidationError
from process_swimlane.workbook import HEADERS, effective_value, frequency_value, read_workbook, text_value


class WorkbookTests(unittest.TestCase):
    def test_excel_saved_fixture(self):
        tasks = read_workbook(Path(__file__).parent / "fixtures" / "excel_saved_formulas.xlsx")
        self.assertEqual([task.task_id for task in tasks], ["001", "PR-01"])
        self.assertEqual(tasks[0].description, "Formula task")
        self.assertEqual(tasks[0].annual_frequency, Decimal("0.5"))
        self.assertEqual(tasks[0].next_task_id, "PR-01")
        self.assertIsNone(tasks[1].next_task_id)

    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "input.xlsx"
        self.book = Workbook()
        self.sheet = self.book.active
        self.sheet.title = "Process"
        self.sheet.append(HEADERS)
        self.sheet.append(["001", "Area", "Do a task", "Dept", 0.5, None])

    def read(self):
        self.book.save(self.path)
        return read_workbook(self.path)

    def test_exact_ids_and_source_row(self):
        self.sheet.insert_rows(2)
        task = self.read()[0]
        self.assertEqual((task.row, task.task_id, task.annual_frequency), (3, "001", Decimal("0.5")))

    def test_ignored_content(self):
        self.book.create_sheet("Notes", 0)["A1"] = "=1/0"
        self.sheet["G2"] = "=1/0"
        self.sheet["G20"] = "notes only"
        self.sheet["A25"].number_format = "@"
        self.assertEqual(len(self.read()), 1)

    def test_missing_sheet(self):
        self.sheet.title = "Other"
        with self.assertRaisesRegex(ValidationError, "Process.*missing"):
            self.read()

    def test_headers(self):
        for value in (None, "wrong", "=1"):
            with self.subTest(value=value):
                self.sheet["D1"] = value
                with self.assertRaisesRegex(ValidationError, "D1.*department"):
                    self.read()

    def test_merged_input(self):
        self.sheet.merge_cells("F3:G3")
        with self.assertRaisesRegex(ValidationError, "Merged"):
            self.read()

    def test_partial_row(self):
        self.sheet["D4"] = "Dept"
        with self.assertRaisesRegex(ValidationError, "A4"):
            self.read()

    def test_numeric_id(self):
        self.sheet["A2"] = 1
        self.sheet["A2"].number_format = "000"
        with self.assertRaisesRegex(ValidationError, "A2.*stored as text"):
            self.read()

    def test_missing_formula_value(self):
        self.sheet["A4"] = '=TEXT(1,"000")'
        with self.assertRaisesRegex(ValidationError, "A4.*Recalculate and save"):
            self.read()

    def test_formula_reader_values(self):
        cached = Workbook().active
        self.sheet["A2"] = '=TEXT(1,"000")'
        for value in ("001", 1, 0, ""):
            with self.subTest(value=value):
                cached["A2"] = value
                result = effective_value(self.sheet["A2"], cached["A2"])
                self.assertEqual(result, value)
        for value in (None, "#VALUE!"):
            cached["A2"] = value
            with self.assertRaisesRegex(ValidationError, "Recalculate and save"):
                effective_value(self.sheet["A2"], cached["A2"])

    def test_values(self):
        for value in (None, ""):
            self.assertIsNone(frequency_value(value, 2))
            self.assertIsNone(text_value(value, 2, "F", "next_task_id", True))
        for value in (0, 0.5, "0.5", "10", Decimal("0.25")):
            self.assertEqual(frequency_value(value, 2), Decimal(str(value)))
        for value in (-1, "-1", True, float("nan"), float("inf"), "0,5", "NaN", "word"):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                frequency_value(value, 2)
        self.assertEqual(text_value(" PR-01 ", 2, "A", "task_id"), " PR-01 ")
        for value in (1, True, None, "", "  "):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                text_value(value, 2, "A", "task_id")
