"""Rebuild the synthetic MVP example workbooks with openpyxl."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from process_swimlane.workbook import HEADERS

STEPS = [
    ("Request", "Identify the required supplies", "Requesting"),
    ("Request", "Describe quantities and purpose", "Requesting"),
    ("Review", "Check the request is complete", "Purchasing"),
    ("Review", "Confirm the available budget", "Finance"),
    ("Order", "Request supplier quotations", "Purchasing"),
    ("Order", "Select a suitable quotation", "Purchasing"),
    ("Order", "Approve the proposed spend", "Finance"),
    ("Order", "Issue the purchase order", "Purchasing"),
    ("Delivery", "Receive the ordered supplies", "Receiving"),
    ("Delivery", "Check delivered quantities", "Receiving"),
    ("Close", "Confirm the supplies meet the request", "Requesting"),
    ("Close", "Record completion of the request", "Purchasing"),
]


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "Process"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    for index, (section, description, department) in enumerate(STEPS, 1):
        sheet.append([f"{index:03}", section, description, department, [120, 0.5, None, 0][(index-1) % 4], f"{index+1:03}" if index < len(STEPS) else None])
    for column, width in zip("ABCDEF", (16, 18, 50, 22, 22, 18)):
        sheet.column_dimensions[column].width = width
    for column in ("A", "F"):
        sheet.column_dimensions[column].number_format = "@"
        for row in range(2, 14):
            sheet[f"{column}{row}"].number_format = "@"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24465C")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    for row in sheet.iter_rows(min_row=2, max_row=13):
        sheet.row_dimensions[row[0].row].height = 34
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EFF5F8")
    sheet.auto_filter.ref = "A1:F13"
    return book


if __name__ == "__main__":
    directory = Path("examples")
    directory.mkdir(exist_ok=True)
    book = workbook()
    book.save(directory / "purchase_request.xlsx")
    for name, cell, value in [
        ("invalid_missing_department", "D1", None),
        ("invalid_multiple_successors", "F2", "002,003"),
        ("invalid_loop", "F13", "001"),
    ]:
        book = workbook()
        book["Process"][cell] = value
        book.save(directory / f"{name}.xlsx")
