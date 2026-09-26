from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from process_swimlane.workbook import read_workbook
from process_swimlane.validation import validate_process

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


class ExampleTests(unittest.TestCase):
    def test_template(self):
        book = load_workbook(EXAMPLES / "purchase_request.xlsx")
        sheet = book["Process"]
        self.assertEqual(sheet.max_row, 13)
        for row in range(2, 14):
            for col in ("A", "F"):
                self.assertEqual(sheet[f"{col}{row}"].number_format, "@")
        book.close()
        tasks = validate_process(read_workbook(EXAMPLES / "purchase_request.xlsx"))
        self.assertEqual(len(tasks), 12)
        self.assertEqual(len(set(t.department for t in tasks)), 4)

    def test_demo_repeat_and_invalid_preservation(self):
        source = EXAMPLES / "purchase_request.xlsx"
        original = source.read_bytes()
        with TemporaryDirectory() as directory:
            def run(path):
                return subprocess.run([sys.executable, "-m", "process_swimlane.cli", str(path)], cwd=directory, capture_output=True, text=True)
            for _ in range(2):
                result = run(source)
                self.assertEqual(result.returncode, 0, result.stderr)
            output = Path(directory) / "output"
            self.assertEqual(sorted(p.name for p in output.iterdir()), ["purchase_request_swimlane.svg", "purchase_request_swimlane_2.svg"])
            before = {p.name:p.read_bytes() for p in output.iterdir()}
            root = ET.fromstring(before["purchase_request_swimlane.svg"])
            for kind, count in (("task", 12), ("lane", 4), ("connection", 11)):
                self.assertEqual(len(root.findall(f".//*[@data-kind='{kind}']")), count)
            for invalid in EXAMPLES.glob("invalid_*.xlsx"):
                result = run(invalid)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("Error:", result.stderr)
                self.assertEqual(before, {p.name:p.read_bytes() for p in output.iterdir()})
        self.assertEqual(source.read_bytes(), original)
